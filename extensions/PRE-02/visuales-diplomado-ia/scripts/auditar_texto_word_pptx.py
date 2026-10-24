#!/usr/bin/env python3
"""Audita fidelidad textual estricta entre un DOCX y un PPTX por diapositiva.

El DOCX debe separar cada lámina con un párrafo del tipo ``DIAPOSITIVA 1``.
El marcador funciona como control estructural y no se considera contenido visible.
El proceso devuelve 0 solo cuando estructura e inventario textual coinciden y cualquier
imagen u objeto complejo fue revisado expresamente. La secuencia geométrica también se exige,
salvo que se active de forma explícita el modo de composición no lineal.
"""

from __future__ import annotations

import argparse
from collections import Counter
import difflib
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from pptx import Presentation
except ImportError as exc:  # pragma: no cover - depende del entorno de destino
    raise SystemExit(
        "Faltan dependencias. Ejecuta el auditor en el entorno de presentaciones de Codex "
        "o instala python-docx y python-pptx. Detalle: " + str(exc)
    )


MARKER_RE = re.compile(
    r"^\s*(?:#+\s*)?DIAPOSITIVA\s+0*(\d+)"
    r"(?:\s*[:\-–—]\s*(.*?))?\s*$",
    re.IGNORECASE,
)
TOKEN_RE = re.compile(r"\w+(?:[’'\-]\w+)*|[^\w\s]", re.UNICODE)


@dataclass
class WordExtraction:
    slides: dict[int, list[str]] = field(default_factory=dict)
    preamble: list[str] = field(default_factory=list)
    duplicate_markers: list[int] = field(default_factory=list)


@dataclass
class SlideExtraction:
    text: str
    pictures: list[str] = field(default_factory=list)
    complex_objects: list[str] = field(default_factory=list)


@dataclass
class SlideResult:
    number: int
    expected: str
    actual: str
    expected_tokens: list[str]
    actual_tokens: list[str]
    missing: list[str]
    extra: list[str]
    sequence_equal: bool
    inventory_equal: bool
    nonlinear_layout_allowed: bool = False

    @property
    def passed(self) -> bool:
        sequence_ok = self.sequence_equal or (
            self.nonlinear_layout_allowed and self.inventory_equal
        )
        return not self.missing and not self.extra and sequence_ok


def clean_text(value: str) -> str:
    """Normaliza solo diferencias de maquetación, nunca ortografía o puntuación."""
    value = unicodedata.normalize("NFC", value)
    value = value.replace("\u00a0", " ").replace("\u00ad", "")
    return re.sub(r"\s+", " ", value).strip()


def tokens(value: str) -> list[str]:
    return TOKEN_RE.findall(clean_text(value))


def iter_docx_blocks(document) -> Iterable[Paragraph | Table]:
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def iter_table_paragraphs(table: Table) -> Iterable[str]:
    for row in table.rows:
        seen_cells: set[int] = set()
        for cell in row.cells:
            cell_key = id(cell._tc)
            if cell_key in seen_cells:
                continue
            seen_cells.add(cell_key)
            for paragraph in cell.paragraphs:
                text = clean_text(paragraph.text)
                if text:
                    yield text


def extract_word(path: Path) -> WordExtraction:
    document = Document(path)
    result = WordExtraction()
    current_slide: int | None = None

    for block in iter_docx_blocks(document):
        paragraphs = [block.text] if isinstance(block, Paragraph) else list(iter_table_paragraphs(block))
        for raw_text in paragraphs:
            text = clean_text(raw_text)
            if not text:
                continue
            marker = MARKER_RE.match(text)
            if marker:
                number = int(marker.group(1))
                if number in result.slides:
                    result.duplicate_markers.append(number)
                result.slides.setdefault(number, [])
                current_slide = number
                inline_title = clean_text(marker.group(2) or "")
                if inline_title:
                    result.slides[number].append(inline_title)
                continue
            if current_slide is None:
                result.preamble.append(text)
            else:
                result.slides[current_slide].append(text)

    return result


def shape_type_name(shape) -> str:
    shape_type = getattr(shape, "shape_type", None)
    return getattr(shape_type, "name", str(shape_type))


def extract_paragraph_texts(text_frame) -> list[str]:
    return [clean_text(p.text) for p in text_frame.paragraphs if clean_text(p.text)]


def extract_slide(slide) -> SlideExtraction:
    blocks: list[tuple[int, int, int, str]] = []
    pictures: list[str] = []
    complex_objects: list[str] = []
    sequence = 0

    def add(text: str, top: int, left: int) -> None:
        nonlocal sequence
        text = clean_text(text)
        if text:
            blocks.append((top, left, sequence, text))
            sequence += 1

    def visit(shape, parent_top: int = 0, parent_left: int = 0) -> None:
        top = parent_top + int(getattr(shape, "top", 0) or 0)
        left = parent_left + int(getattr(shape, "left", 0) or 0)
        type_name = shape_type_name(shape)
        name = getattr(shape, "name", type_name)

        if type_name == "GROUP":
            for child in shape.shapes:
                visit(child, top, left)
            return

        if type_name in {"PICTURE", "LINKED_PICTURE"}:
            pictures.append(name)

        has_table = bool(getattr(shape, "has_table", False))
        has_chart = bool(getattr(shape, "has_chart", False))

        if has_table:
            table = shape.table
            for row_index, row in enumerate(table.rows):
                for col_index, cell in enumerate(row.cells):
                    for paragraph_index, text in enumerate(extract_paragraph_texts(cell.text_frame)):
                        add(
                            text,
                            top + row_index * 10_000 + paragraph_index,
                            left + col_index * 10_000,
                        )
            return

        if has_chart:
            complex_objects.append(f"{name} (gráfica)")
            return

        if type_name in {"SMART_ART", "DIAGRAM", "IGX_GRAPHIC"}:
            complex_objects.append(f"{name} ({type_name})")
            return

        if type_name == "GRAPHIC_FRAME" and not has_table:
            complex_objects.append(f"{name} (objeto gráfico no tabular)")
            return

        if bool(getattr(shape, "has_text_frame", False)):
            for paragraph_index, text in enumerate(extract_paragraph_texts(shape.text_frame)):
                add(text, top + paragraph_index, left)

    for shape in slide.shapes:
        visit(shape)

    blocks.sort(key=lambda item: (item[0], item[1], item[2]))
    return SlideExtraction(
        text=clean_text(" ".join(item[3] for item in blocks)),
        pictures=pictures,
        complex_objects=complex_objects,
    )


def diff_tokens(expected: list[str], actual: list[str]) -> tuple[list[str], list[str]]:
    missing: list[str] = []
    extra: list[str] = []
    matcher = difflib.SequenceMatcher(a=expected, b=actual, autojunk=False)
    for opcode, a1, a2, b1, b2 in matcher.get_opcodes():
        if opcode in {"delete", "replace"} and a1 != a2:
            missing.append(" ".join(expected[a1:a2]))
        if opcode in {"insert", "replace"} and b1 != b2:
            extra.append(" ".join(actual[b1:b2]))
    return missing, extra


def counter_diff_tokens(expected: list[str], actual: list[str]) -> tuple[list[str], list[str]]:
    """Detecta faltantes y sobrantes reales sin confundirlos con reordenación espacial."""
    expected_counter = Counter(expected)
    actual_counter = Counter(actual)
    missing = list((expected_counter - actual_counter).elements())
    extra = list((actual_counter - expected_counter).elements())
    return missing, extra


def compare_slide(
    number: int,
    expected: str,
    actual: str,
    allow_nonlinear_layout: bool = False,
) -> SlideResult:
    expected_tokens = tokens(expected)
    actual_tokens = tokens(actual)
    inventory_equal = Counter(expected_tokens) == Counter(actual_tokens)
    if allow_nonlinear_layout:
        missing, extra = counter_diff_tokens(expected_tokens, actual_tokens)
    else:
        missing, extra = diff_tokens(expected_tokens, actual_tokens)
    return SlideResult(
        number=number,
        expected=expected,
        actual=actual,
        expected_tokens=expected_tokens,
        actual_tokens=actual_tokens,
        missing=missing,
        extra=extra,
        sequence_equal=expected_tokens == actual_tokens,
        inventory_equal=inventory_equal,
        nonlinear_layout_allowed=allow_nonlinear_layout,
    )


def shortened(value: str, limit: int = 500) -> str:
    return value if len(value) <= limit else value[: limit - 1] + "…"


def build_report(
    word_path: Path,
    pptx_path: Path,
    word: WordExtraction,
    ppt_slides: list[SlideExtraction],
    results: list[SlideResult],
    structural_errors: list[str],
    image_review_pending: list[int],
    complex_review_pending: list[int],
    authorized_transformations: list[str] | None = None,
) -> str:
    passed = (
        not structural_errors
        and all(result.passed for result in results)
        and not image_review_pending
        and not complex_review_pending
    )
    lines = [
        "# Auditoría hostil Word–PowerPoint",
        "",
        f"- Word: `{word_path}`",
        f"- PowerPoint: `{pptx_path}`",
        f"- Resultado: **{'APROBADO' if passed else 'RECHAZADO'}**",
        f"- Secciones Word: {len(word.slides)}",
        f"- Diapositivas PowerPoint: {len(ppt_slides)}",
        "",
    ]

    if structural_errors:
        lines.extend(["## Errores estructurales", ""])
        lines.extend(f"- {error}" for error in structural_errors)
        lines.append("")

    if authorized_transformations:
        lines.extend(["## Transformaciones expresamente autorizadas", ""])
        lines.extend(f"- {item}" for item in authorized_transformations)
        lines.append("")

    if image_review_pending:
        lines.extend(
            [
                "## Revisión raster pendiente",
                "",
                "Las siguientes diapositivas contienen imágenes. No pueden aprobarse hasta "
                "confirmar visualmente que no incluyen texto rasterizado:",
                "",
                "- " + ", ".join(map(str, image_review_pending)),
                "",
            ]
        )

    if complex_review_pending:
        lines.extend(
            [
                "## Objetos complejos pendientes",
                "",
                "Las siguientes diapositivas contienen gráficas, SmartArt u objetos cuyo "
                "texto requiere comparación manual:",
                "",
                "- " + ", ".join(map(str, complex_review_pending)),
                "",
            ]
        )

    lines.extend(["## Resultado por diapositiva", ""])
    for result in results:
        status = "APROBADA" if result.passed else "RECHAZADA"
        lines.extend(
            [
                f"### Diapositiva {result.number}: {status}",
                "",
                f"- Tokens esperados: {len(result.expected_tokens)}",
                f"- Tokens encontrados: {len(result.actual_tokens)}",
            ]
        )
        if result.missing:
            lines.append("- Falta: " + " | ".join(f"`{shortened(x, 160)}`" for x in result.missing))
        if result.extra:
            lines.append("- Sobra: " + " | ".join(f"`{shortened(x, 160)}`" for x in result.extra))
        if not result.sequence_equal and result.inventory_equal:
            if result.nonlinear_layout_allowed:
                lines.append(
                    "- Inventario literal exacto; el orden espacial difiere por la composición "
                    "no lineal autorizada."
                )
            else:
                lines.append("- El inventario coincide, pero la secuencia textual cambió.")
        if not result.passed:
            lines.extend(
                [
                    "- Esperado: `" + shortened(result.expected).replace("`", "´") + "`",
                    "- Encontrado: `" + shortened(result.actual).replace("`", "´") + "`",
                ]
            )
        lines.append("")

    criterion = (
        "Solo se normalizan saltos de línea, espacios de maquetación, espacio no separable, "
        "guion blando y composición Unicode. Se conservan y comparan palabras, cifras, "
        "mayúsculas, acentos y signos de puntuación."
    )
    if authorized_transformations:
        criterion += (
            " Las únicas excepciones son las transformaciones identificadas arriba y activadas "
            "mediante opciones explícitas del auditor."
        )
    lines.extend(["## Criterio", "", criterion, ""])
    return "\n".join(lines)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Compara literalmente el texto de un DOCX y un PPTX, diapositiva por diapositiva."
    )
    parser.add_argument("word", type=Path, help="Archivo Word .docx con marcadores DIAPOSITIVA N")
    parser.add_argument("powerpoint", type=Path, help="Archivo PowerPoint .pptx que se auditará")
    parser.add_argument("--report", type=Path, help="Ruta opcional para guardar el informe Markdown")
    parser.add_argument(
        "--confirm-images-without-text",
        action="store_true",
        help="Usar solo después de revisar visualmente que ninguna imagen contiene texto.",
    )
    parser.add_argument(
        "--confirm-complex-objects-audited",
        action="store_true",
        help="Usar solo después de comparar manualmente el texto de gráficas o SmartArt.",
    )
    parser.add_argument(
        "--assign-preamble-to-slide-one",
        action="store_true",
        help=(
            "Asigna a la diapositiva 1 el texto anterior al primer marcador. "
            "Úsalo solo cuando el usuario haya confirmado ese destino."
        ),
    )
    parser.add_argument(
        "--uppercase-titles",
        action="store_true",
        help=(
            "Convierte a mayúsculas el primer bloque de cada sección Word antes de comparar. "
            "Úsalo solo cuando el usuario haya pedido títulos en mayúsculas."
        ),
    )
    parser.add_argument(
        "--allow-nonlinear-layout",
        action="store_true",
        help=(
            "Compara el inventario literal de tokens con multiplicidad, sin imponer lectura "
            "geométrica superior-izquierda. Úsalo solo en escenas ascendentes, circulares, "
            "bilaterales o con soportes distribuidos; cualquier token faltante o sobrante "
            "continúa provocando rechazo."
        ),
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if not args.word.is_file() or args.word.suffix.lower() != ".docx":
        print("ERROR: el primer archivo debe ser un DOCX existente.", file=sys.stderr)
        return 2
    if not args.powerpoint.is_file() or args.powerpoint.suffix.lower() != ".pptx":
        print("ERROR: el segundo archivo debe ser un PPTX existente.", file=sys.stderr)
        return 2

    word = extract_word(args.word)
    authorized_transformations: list[str] = []
    if args.uppercase_titles:
        for slide_blocks in word.slides.values():
            if slide_blocks:
                slide_blocks[0] = slide_blocks[0].upper()
        authorized_transformations.append(
            "Los títulos, identificados como el primer bloque de cada sección, se comparan en MAYÚSCULAS."
        )
    if args.assign_preamble_to_slide_one and word.preamble:
        word.slides[1] = list(word.preamble) + word.slides.get(1, [])
        word.preamble.clear()
        authorized_transformations.append(
            "El texto anterior a DIAPOSITIVA 1 se asigna, en el mismo orden, a la diapositiva 1."
        )
    if args.allow_nonlinear_layout:
        authorized_transformations.append(
            "La composición no lineal puede alterar el orden geométrico de lectura; se exige "
            "inventario literal idéntico, con la misma multiplicidad de palabras, cifras y signos."
        )
    presentation = Presentation(args.powerpoint)
    ppt_slides = [extract_slide(slide) for slide in presentation.slides]

    structural_errors: list[str] = []
    if not word.slides:
        structural_errors.append(
            "El Word no contiene marcadores del tipo DIAPOSITIVA N; no existe una asignación "
            "determinista por lámina."
        )
    if word.preamble:
        structural_errors.append(
            "Existe texto antes del primer marcador sin destino de diapositiva: "
            + " | ".join(word.preamble)
        )
    if word.duplicate_markers:
        structural_errors.append(
            "Marcadores duplicados: " + ", ".join(map(str, sorted(set(word.duplicate_markers))))
        )

    expected_numbers = set(range(1, len(ppt_slides) + 1))
    source_numbers = set(word.slides)
    missing_sections = sorted(expected_numbers - source_numbers)
    out_of_range = sorted(source_numbers - expected_numbers)
    if missing_sections:
        structural_errors.append(
            "No existe sección Word para las diapositivas: " + ", ".join(map(str, missing_sections))
        )
    if out_of_range:
        structural_errors.append(
            "Existen secciones Word sin diapositiva PowerPoint: " + ", ".join(map(str, out_of_range))
        )
    if len(word.slides) != len(ppt_slides):
        structural_errors.append(
            f"La cantidad no coincide: {len(word.slides)} secciones Word frente a "
            f"{len(ppt_slides)} diapositivas PowerPoint."
        )

    results: list[SlideResult] = []
    for number in sorted(expected_numbers | source_numbers):
        expected = clean_text(" ".join(word.slides.get(number, [])))
        actual = ppt_slides[number - 1].text if 1 <= number <= len(ppt_slides) else ""
        results.append(
            compare_slide(
                number,
                expected,
                actual,
                allow_nonlinear_layout=args.allow_nonlinear_layout,
            )
        )

    image_slides = [i for i, slide in enumerate(ppt_slides, start=1) if slide.pictures]
    complex_slides = [i for i, slide in enumerate(ppt_slides, start=1) if slide.complex_objects]
    image_review_pending = [] if args.confirm_images_without_text else image_slides
    complex_review_pending = [] if args.confirm_complex_objects_audited else complex_slides

    report = build_report(
        args.word,
        args.powerpoint,
        word,
        ppt_slides,
        results,
        structural_errors,
        image_review_pending,
        complex_review_pending,
        authorized_transformations,
    )
    print(report)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(report, encoding="utf-8")

    approved = (
        not structural_errors
        and all(result.passed for result in results)
        and not image_review_pending
        and not complex_review_pending
    )
    return 0 if approved else 1


if __name__ == "__main__":
    raise SystemExit(main())
